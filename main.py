import subprocess
import json
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from datetime import datetime

moeda = input("Digite a moeda (ex: USD-BRL): ")
numero_dias = input("Digite o número de dias: ")

resultado = subprocess.run(
    ["node", "api.js", moeda, numero_dias],
    capture_output=True,
    text=True
)

if resultado.returncode != 0:
    print("Erro:", resultado.stderr)
    exit()

dados = json.loads(resultado.stdout)

# Criar Excel
wb = Workbook()

ws = wb.active
ws.title = "Dados"

# Cabeçalho
ws.append([
    "Data",
    "Moeda",
    "Máxima",
    "Mínima",
    "Variação",
    "Variação %",
    "Compra",
    "Venda"
])

# Inserir dados
for item in dados:
    data = datetime.fromtimestamp(int(item["timestamp"]))

    ws.append([
        data,
        item.get("code", moeda.split("-")[0]),
        float(item["high"]),
        float(item["low"]),
        float(item["varBid"]),
        float(item["pctChange"]),
        float(item["bid"]),
        float(item["ask"])
    ])

# Criar gráfico da cotação
grafico_cotacao = LineChart()

grafico_cotacao.title = "Cotação USD/BRL"
grafico_cotacao.y_axis.title = "Valor (R$)"
grafico_cotacao.x_axis.title = "Data"

valores = Reference(ws, min_col=7, min_row=1, max_row=ws.max_row)
datas = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)

grafico_cotacao.add_data(valores, titles_from_data=True)
grafico_cotacao.set_categories(datas)

grafico_cotacao.height = 10
grafico_cotacao.width = 20

ws.add_chart(grafico_cotacao, "J2")


# Criar gráfico de variação percentual
grafico_variacao = LineChart()

grafico_variacao.title = "Variação percentual"
grafico_variacao.y_axis.title = "Variação (%)"
grafico_variacao.x_axis.title = "Data"

variacao = Reference(ws, min_col=6, min_row=1, max_row=ws.max_row)

grafico_variacao.add_data(variacao, titles_from_data=True)
grafico_variacao.set_categories(datas)

grafico_variacao.height = 10
grafico_variacao.width = 20

ws.add_chart(grafico_variacao, "J22")


# Salvar
arquivo = "cotacao.xlsx"
wb.save(arquivo)

print(f"Excel criado: {arquivo}")